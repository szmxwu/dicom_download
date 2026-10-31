import gzip,hashlib,json,os,ssl,subprocess,threading,time,urllib.error,urllib.request
from pathlib import Path
import nibabel as nib
import numpy as np
import pytest
from PIL import Image,PngImagePlugin
from openpyxl import Workbook,load_workbook
from src.research_export.publish import publish,digest,Rejected
from src.research_export.gateway import Store,Server
from src.research_export.fetch import Fetch

@pytest.fixture
def source(tmp_path):
 p=tmp_path/'source';p.mkdir();a=np.arange(24,dtype=np.int16).reshape(2,3,4);im=nib.Nifti1Image(a,np.diag([2.,3.,4.,1.]));im.header['descrip']=b'SYNTHETIC_SECRET';im.header['aux_file']=b'SYNTHETIC_SECRET';im.header.extensions.append(nib.nifti1.Nifti1Extension(6,b'SYNTHETIC_SECRET'));nib.save(im,p/'image.nii.gz')
 np.savez_compressed(p/'image.npz',data=a.astype(np.float32))
 info=PngImagePlugin.PngInfo();info.add_text('PatientID','SYNTHETIC_SECRET');Image.fromarray(np.arange(64,dtype=np.uint8).reshape(8,8)).save(p/'preview.png',pnginfo=info)
 book=Workbook();book.active.title='DICOM_Metadata';book.active.append(['PatientID','AccessionNumber','StudyDate','FileName','Rows','Modality','SeriesNumber']);book.active.append(['SYNTHETIC_SECRET','SYNTHETIC_SECRET','20000101','SYNTHETIC_SECRET',8,'CT',1]);sheet=book.create_sheet('Series_Summary');sheet.append(['PatientID','AccessionNumber','StudyDate','Rows']);sheet.append(['SYNTHETIC_SECRET','SYNTHETIC_SECRET','20000101',8]);book.save(p/'metadata.xlsx')
 review={'policy_version':'research-v1','reviewer':'synthetic-test','files':[{'path':x.name,'case_ref':'SYNTHETIC_CASE','series_number':1,'sha256':digest(x),'pixel_review':{'status':'approved','evidence':'synthetic-array-fixture'},'quality_review':{'status':'approved','evidence':'synthetic-values-and-geometry'}} for x in sorted(p.iterdir())]}
 return p,review

def test_publisher_clean_all_formats(source,tmp_path):
 src,review=source;root=tmp_path/'export';m=publish(src,review,root,'synthetic',b'k'*32);base=root/'published'/m['export_id'];assert m['case_count']==1
 for a in m['artifacts']:
  p=base/'artifacts'/a['file'];assert digest(p)==a['sha256'];assert not p.stat().st_mode&0o222;assert 'SYNTHETIC' not in a['name']
  if a['format']=='.nii.gz':
   im=nib.load(p);assert bytes(im.header['descrip']).strip(b'\x00')==b'';assert len(im.header.extensions)==0;np.testing.assert_array_equal(im.get_fdata(),nib.load(src/'image.nii.gz').get_fdata());np.testing.assert_array_equal(im.affine,nib.load(src/'image.nii.gz').affine)
  elif a['format']=='.xlsx':
   book=load_workbook(p)
   for sheet in book:
    values=list(sheet.values);assert 'ResearchID' in values[0];assert all('SYNTHETIC_SECRET' not in str(row) and 'StudyDate' not in str(row) for row in values)
  elif a['format']=='.png':assert not Image.open(p).info
  else:
   with np.load(p,allow_pickle=False) as z:np.testing.assert_array_equal(z['data'],np.load(src/'image.npz')['data'])

@pytest.mark.parametrize('change',['review','digest','pixel','unknown','symlink','npz_object','multi_case_excel','trailing_nifti'])
def test_publisher_rejects(source,tmp_path,change):
 src,review=source
 if change=='review':review['reviewer']=''
 elif change=='digest':review['files'][0]['sha256']='0'*64
 elif change=='pixel':
  for x in review['files']:x['pixel_review']={}
 elif change=='unknown':(src/'x.dcm').write_bytes(b'x');review['files']=[{'path':'x.dcm','case_ref':'x','sha256':digest(src/'x.dcm')}]
 elif change=='symlink':
  (src/'alias.png').symlink_to(src/'preview.png');review['files']=[{'path':'alias.png','case_ref':'x','sha256':digest(src/'preview.png'),'pixel_review':{'status':'approved','evidence':'test'}}]
 elif change=='npz_object':
  np.savez(src/'image.npz',data=np.array(['SYNTHETIC_SECRET'],dtype=object))
  for x in review['files']:
   if x['path']=='image.npz':x['sha256']=digest(src/x['path'])
 elif change=='multi_case_excel':
  book=load_workbook(src/'metadata.xlsx');book.active.append(['OTHER_CASE','OTHER_ACC','20000101','x',8,'CT',2]);book.save(src/'metadata.xlsx')
  for x in review['files']:
   if x['path']=='metadata.xlsx':x['sha256']=digest(src/x['path'])
 else:
  with gzip.open(src/'image.nii.gz','ab') as f:f.write(b'SYNTHETIC_SECRET')
  for x in review['files']:
   if x['path']=='image.nii.gz':x['sha256']=digest(src/x['path'])
 with pytest.raises((Rejected,ValueError)):publish(src,review,tmp_path/'export','test',b'k'*32)
 assert not list((tmp_path/'export'/'published').glob('*'))

@pytest.fixture(scope='module')
def certificates(tmp_path_factory):
 p=tmp_path_factory.mktemp('certs')
 def run(*a):subprocess.run(['openssl',*a],cwd=p,check=True,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
 run('req','-x509','-newkey','rsa:2048','-nodes','-keyout','ca.key','-out','ca.crt','-subj','/CN=synthetic-export-ca','-days','1','-addext','basicConstraints=critical,CA:TRUE','-addext','keyUsage=critical,keyCertSign,cRLSign')
 for name,purpose in [('server','serverAuth'),('client','clientAuth'),('other','clientAuth')]:
  run('req','-new','-newkey','rsa:2048','-nodes','-keyout',name+'.key','-out',name+'.csr','-subj','/CN='+name)
  (p/(name+'.ext')).write_text('basicConstraints=critical,CA:FALSE\nkeyUsage=critical,digitalSignature,keyEncipherment\nextendedKeyUsage='+purpose+'\nsubjectAltName=IP:127.0.0.1\nauthorityKeyIdentifier=keyid,issuer\nsubjectKeyIdentifier=hash\n')
  run('x509','-req','-in',name+'.csr','-CA','ca.crt','-CAkey','ca.key','-CAcreateserial','-out',name+'.crt','-days','1','-extfile',name+'.ext')
 return p

@pytest.fixture
def gateway(source,tmp_path,certificates):
 src,review=source;root=tmp_path/'export';m=publish(src,review,root,'test',b'k'*32);gid='a'*32;p=certificates;fp=hashlib.sha256(ssl.PEM_cert_to_DER_cert((p/'client.crt').read_text())).hexdigest();policy=tmp_path/'policy.json';g={'status':'APPROVED','certificate_sha256':fp,'export_id':m['export_id'],'project_id':'test','manifest_sha256':digest(root/'published'/m['export_id']/'manifest.json'),'not_before':0,'expires_at':time.time()+600,'max_cases':100,'max_bytes':1000000,'rate_bytes_s':1000000,'concurrency':1};policy.write_text(json.dumps({'grants':{gid:g}}));store=Store(policy,root/'published',tmp_path/'state');ctx=ssl.SSLContext(ssl.PROTOCOL_TLS_SERVER);ctx.verify_mode=ssl.CERT_REQUIRED;ctx.load_cert_chain(p/'server.crt',p/'server.key');ctx.load_verify_locations(p/'ca.crt');server=Server(('127.0.0.1',0),store,ctx);thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
 def get(path,headers=None,client='client',method='GET'):
  c=ssl.create_default_context(cafile=p/'ca.crt')
  if client:c.load_cert_chain(p/(client+'.crt'),p/(client+'.key'))
  opener=urllib.request.build_opener(urllib.request.ProxyHandler({}),urllib.request.HTTPSHandler(context=c));return opener.open(urllib.request.Request('https://127.0.0.1:'+str(server.server_port)+path,headers={'X-Export-Grant':gid,**(headers or {})},method=method),timeout=4)
 yield get,m,gid,policy,g,store
 server.shutdown();server.server_close();thread.join(timeout=3)

def test_manifest_and_range(gateway):
 get,m,gid,_,_,_=gateway
 with get('/v1/exports/'+m['export_id']+'/manifest') as r:assert all('file' not in a for a in json.load(r)['artifacts'])
 a=m['artifacts'][0];url='/v1/artifacts/'+a['artifact_id']+'/content'
 with get(url) as r:data=r.read();assert hashlib.sha256(data).hexdigest()==a['sha256']
 with get(url,{'Range':'bytes=1-3','If-Match':'"'+a['sha256']+'"'}) as r:assert r.status==206 and r.read()==data[1:4]
 for headers,status in [({'Range':'bytes=1-3'},416),({'If-Match':'"bad"'},412),({'Range':'bytes=0-1,2-3','If-Match':'"'+a['sha256']+'"'},416)]:
  with pytest.raises(urllib.error.HTTPError) as e:get(url,headers)
  assert e.value.code==status

@pytest.mark.parametrize('path',['/api/pacs-config','/api/cache/clear','/api/process/single','/api/tasks/history','/v1/artifacts/'+('0'*32)+'/content','/v1/exports/'+('0'*32)+'/manifest','/v1/exports/../manifest','/v1/artifacts/%2e%2e/content'])
def test_no_legacy_or_cross_scope(gateway,path):
 get,*_=gateway
 with pytest.raises(urllib.error.HTTPError) as e:get(path)
 assert e.value.code==404

def test_certificates_revocation_expiry_quota(gateway):
 get,m,gid,policy,g,store=gateway;url='/v1/exports/'+m['export_id']+'/manifest'
 with pytest.raises(Exception):get(url,client=None)
 with pytest.raises(urllib.error.HTTPError) as e:get(url,client='other')
 assert e.value.code==404
 for updates in ({'status':'REVOKED'},{'expires_at':1},{'max_bytes':1}):
  new=dict(g,**updates);policy.write_text(json.dumps({'grants':{gid:new}}))
  with pytest.raises(urllib.error.HTTPError) as e:get(url)
  assert e.value.code==404

def test_active_stream_rechecks_scope(gateway):
 _,_,gid,policy,g,store=gateway;peer=g['certificate_sha256'];store.charge(gid,peer,2,expected=g)
 changed=dict(g,export_id='0'*32);policy.write_text(json.dumps({'grants':{gid:changed}}))
 from src.research_export.gateway import Denied
 with pytest.raises(Denied):store.charge(gid,peer,2,expected=g)

def test_fetch_rejects_arbitrary_network_parameters(tmp_path):
 config=tmp_path/'config.json';config.write_text(json.dumps({'server':'https://192.0.0.193:5005','grants':{}}));f=Fetch(config)
 with pytest.raises(ValueError):f.settings('a'*32)

def test_revoke_interrupts_stream_and_quota_counts_resends(gateway):
 get,m,gid,policy,g,store=gateway
 # A synthetic larger payload exercises streamed checks independently of publisher tests.
 a=m['artifacts'][0];path=store.published/m['export_id']/'artifacts'/a['file'];os.chmod(path,0o600);payload=os.urandom(1024*1024);path.write_bytes(payload);os.chmod(path,0o440);a['bytes']=len(payload);a['sha256']=hashlib.sha256(payload).hexdigest()
 manifest=store.published/m['export_id']/'manifest.json';os.chmod(manifest,0o600);manifest.write_text(json.dumps(m));os.chmod(manifest,0o440);g=dict(g,manifest_sha256=digest(manifest),max_bytes=3*1024**2,rate_bytes_s=64*1024);policy.write_text(json.dumps({'grants':{gid:g}}));url='/v1/artifacts/'+a['artifact_id']+'/content'
 response=get(url);assert response.read(1024)
 policy.write_text(json.dumps({'grants':{gid:dict(g,status='REVOKED')}}))
 import http.client
 with pytest.raises(http.client.IncompleteRead):response.read()
 response.close()
 with pytest.raises(urllib.error.HTTPError):get(url)
 with store.db() as db:assert db.execute('SELECT bytes FROM usage WHERE grant_id=?',(gid,)).fetchone()[0]>=256*1024

def test_non_read_methods_rejected(gateway):
 get,m,*_=gateway
 for method in ('POST','PUT','DELETE','PATCH','TRACE','OPTIONS'):
  with pytest.raises(urllib.error.HTTPError) as e:get('/v1/exports/'+m['export_id'],method=method)
  assert e.value.code==405

def test_fetch_async_resume_and_principal_scope(gateway,tmp_path,monkeypatch):
 import pwd
 get,m,gid,policy,g,store=gateway;c={'clients':[pwd.getpwuid(os.getuid()).pw_name],'downloads':str(tmp_path/'downloads'),'state':str(tmp_path/'jobs')};Path(c['downloads']).mkdir();Path(c['state']).mkdir();config=tmp_path/'fetch.json';config.write_text(json.dumps(c));f=Fetch(config)
 monkeypatch.setattr(f,'settings',lambda grant:(c,m['export_id'],None))
 monkeypatch.setattr(f,'request',lambda grant,path,headers=None:get(path,headers))
 aid=m['artifacts'][0]['artifact_id'];a=m['artifacts'][0];folder=Path(c['downloads'])/m['export_id'];folder.mkdir();original=(store.published/m['export_id']/'artifacts'/a['file']).read_bytes();(folder/(aid+a['format']+'.partial')).write_bytes(original[:10])
 with pytest.raises(ValueError):f.call({'operation':'download','grant_id':gid,'artifact_id':aid},os.getuid()+10000)
 with pytest.raises(ValueError):f.call({'operation':'download','grant_id':gid,'artifact_id':aid,'url':'http://127.0.0.1:5005'},os.getuid())
 job=f.call({'operation':'download','grant_id':gid,'artifact_id':aid},os.getuid());deadline=time.monotonic()+10
 while f.jobs[job['job_id']]['state']=='RUNNING' and time.monotonic()<deadline:time.sleep(.05)
 result=f.call({'operation':'job','grant_id':gid,'job_id':job['job_id']},os.getuid());assert result['state']=='SUCCEEDED';assert Path(result['path']).read_bytes()==original
 with pytest.raises(ValueError):f.call({'operation':'job','grant_id':'b'*32,'job_id':job['job_id']},os.getuid())

def test_png_16bit_pixels_preserved(tmp_path):
 from src.research_export.publish import clean_png
 a=np.arange(256,dtype=np.uint16).reshape(16,16)*200;source=tmp_path/'src.png';target=tmp_path/'dst.png';Image.fromarray(a).save(source);clean_png(source,target);np.testing.assert_array_equal(np.asarray(Image.open(target)),a)

def test_excel_cross_sheet_case_mismatch_rejected(source,tmp_path):
 src,review=source;book=load_workbook(src/'metadata.xlsx');book['Series_Summary']['A2']='OTHER_SYNTHETIC_CASE';book.save(src/'metadata.xlsx')
 for row in review['files']:
  if row['path']=='metadata.xlsx':row['sha256']=digest(src/'metadata.xlsx')
 with pytest.raises(Rejected):publish(src,review,tmp_path/'export','test',b'k'*32)

@pytest.mark.parametrize('extension',['.nii.gz','.npz','.png'])
def test_constant_image_not_published(tmp_path,extension):
 from src.research_export.publish import clean_nifti,clean_npz,clean_png
 src=tmp_path/('src'+extension);dst=tmp_path/('dst'+extension)
 if extension=='.nii.gz':nib.save(nib.Nifti1Image(np.zeros((2,2,2),dtype=np.int16),np.eye(4)),src);fn=clean_nifti
 elif extension=='.npz':np.savez_compressed(src,data=np.zeros((3,3),dtype=np.float32));fn=clean_npz
 else:Image.fromarray(np.zeros((3,3),dtype=np.uint8)).save(src);fn=clean_png
 with pytest.raises(Rejected):fn(src,dst)

def test_constant_rgb_preview_rejected(tmp_path):
 from src.research_export.publish import clean_png
 src=tmp_path/'red.png';Image.new('RGB',(4,4),(255,0,0)).save(src)
 with pytest.raises(Rejected):clean_png(src,tmp_path/'out.png')

def test_npz_declared_shape_bomb_rejected_before_allocation(tmp_path):
 import io,zipfile
 from src.research_export.publish import clean_npz
 header=io.BytesIO();np.lib.format.write_array_header_1_0(header,{'descr':'|u1','fortran_order':False,'shape':(10**12,)})
 src=tmp_path/'bomb.npz'
 with zipfile.ZipFile(src,'w') as z:z.writestr('data.npy',header.getvalue())
 with pytest.raises(Rejected):clean_npz(src,tmp_path/'out.npz')

@pytest.mark.parametrize('error,expected',[(ValueError('unknown_grant'),'UNKNOWN_GRANT'),(ValueError('SYNTHETIC_PRIVATE_DETAIL'),'REQUEST_FAILED'),(OSError('/private/SYNTHETIC_MARKER'),'SERVICE_UNAVAILABLE'),(urllib.error.HTTPError('https://private/SYNTHETIC_MARKER',404,'internal detail',{},None),'ACCESS_DENIED_OR_NOT_READY')])
def test_fetch_errors_are_bounded_and_redacted(error,expected):
 from src.research_export.fetch import error_code
 assert error_code(error)==expected

# Multi-day queue/lease tests use only synthetic files and an isolated local mTLS server.
@pytest.fixture
def durable_factory(gateway, tmp_path, monkeypatch):
 import pwd, urllib.parse
 from src.research_export.queue import DurableFetch
 get,m,gid,policy,g,store=gateway
 index=[0]
 def make(existing=None):
  if existing is None:
   index[0]+=1;home=tmp_path/str(index[0]);home.mkdir();(home/'downloads').mkdir();(home/'jobs').mkdir()
   config=home/'config.json';config.write_text(json.dumps({'server':'https://192.0.0.193:8443','clients':[pwd.getpwuid(os.getuid()).pw_name],'grants':{gid:m['export_id']},'downloads':str(home/'downloads'),'state':str(home/'jobs'),'min_free_bytes':0}))
  else:config=existing
  f=DurableFetch(config);c=json.loads(config.read_text())
  class Opener:
   def open(self,req,timeout=30):return get(urllib.parse.urlsplit(req.full_url).path,dict(req.header_items()))
  monkeypatch.setattr(f,'settings',lambda grant:(c,m['export_id'],Opener()))
  return f
 return make

def enable_long_grant(gateway):
 get,m,gid,policy,g,store=gateway
 g=dict(g,not_before=int(time.time()),expires_at=int(time.time())+14*86400,lease_seconds=600)
 policy.write_text(json.dumps({'grants':{gid:g}}))
 return g

def test_renewal_keeps_scope_quota_and_hard_deadline(gateway):
 from src.research_export import leases
 from src.research_export.gateway import Denied
 _,m,gid,policy,_,store=gateway;g=enable_long_grant(gateway);peer=g['certificate_sha256']
 a=leases.renew(store,gid,peer);store.charge(gid,peer,123,lease=a['token'])
 b=leases.renew(store,gid,peer);assert a['token']==b['token'];assert 0<b['valid_for']<=600
 with store.db() as db:
  assert db.execute('SELECT bytes FROM usage WHERE grant_id=?',(gid,)).fetchone()[0]==123
  db.execute('UPDATE leases SET until=0 WHERE grant_id=?',(gid,))
 with pytest.raises(Denied):store.charge(gid,peer,1,lease=a['token'])
 # Expired short lease can be renewed within the original approval, not beyond it.
 leases.renew(store,gid,peer);store.charge(gid,peer,1,lease=a['token'])
 policy.write_text(json.dumps({'grants':{gid:dict(g,expires_at=1)}}))
 with pytest.raises(Denied):leases.renew(store,gid,peer)

@pytest.mark.parametrize('change',[{'status':'REVOKED'},{'export_id':'0'*32},{'manifest_sha256':'0'*64},{'expires_at':1}])
def test_lease_cannot_survive_policy_revocation_or_scope_change(gateway,change):
 from src.research_export import leases
 from src.research_export.gateway import Denied
 _,m,gid,policy,_,store=gateway;g=enable_long_grant(gateway);peer=g['certificate_sha256'];token=leases.renew(store,gid,peer)['token']
 policy.write_text(json.dumps({'grants':{gid:dict(g,**change)}}))
 with pytest.raises(Denied):store.charge(gid,peer,1,lease=token)


def test_new_grant_requires_lease_but_has_no_extra_api(gateway):
 get,m,gid,_,_,_=gateway;enable_long_grant(gateway);url='/v1/exports/'+m['export_id']
 with pytest.raises(urllib.error.HTTPError):get(url+'/manifest')
 with get(url+'/lease') as r:lease=json.load(r)
 with get(url+'/manifest',{'X-Export-Lease':lease['token']}) as r:assert r.status==200
 for path in ('/api/process/single',url+'/lease?days=100','/v1/exports/'+('0'*32)+'/lease'):
  with pytest.raises(urllib.error.HTTPError):get(path,{'X-Export-Lease':lease['token']})


def test_queue_batch_is_durable_deduplicated_and_downloads_all(durable_factory,gateway):
 _,m,gid,*_=gateway;enable_long_grant(gateway);f=durable_factory()
 request={'operation':'batch','grant_id':gid};job=f.call(request,os.getuid())
 assert f.call(request,os.getuid())['job_id']==job['job_id']
 # No foreground connection is needed after enqueue; a new service adopts it.
 f=durable_factory(f.config)
 for _ in range(7):f.run_once()
 result=f.call({'operation':'job','grant_id':gid,'job_id':job['job_id']},os.getuid())
 assert result['state']=='SUCCEEDED' and result['completed']==4 and result['total']==4
 c=json.loads(f.config.read_text());files=list((Path(c['downloads'])/m['export_id']).glob('*'))
 assert len(files)==4 and all(not x.name.endswith('partial') for x in files)
 assert all(digest(x)==next(a['sha256'] for a in m['artifacts'] if x.name.startswith(a['artifact_id'])) for x in files)


def test_queue_network_retry_resumes_partial_and_recovers_running(durable_factory,gateway,monkeypatch):
 _,m,gid,*_=gateway;f=durable_factory();a=m['artifacts'][0]
 c=json.loads(f.config.read_text());directory=Path(c['downloads'])/m['export_id'];directory.mkdir()
 original=(gateway[-1].published/m['export_id']/'artifacts'/a['file']).read_bytes()
 partial=directory/(a['artifact_id']+a['format']+'.partial');partial.write_bytes(original[:10])
 job=f.call({'operation':'download','grant_id':gid,'artifact_id':a['artifact_id']},os.getuid())
 real=f.request
 def offline(*args,**kwargs):raise ConnectionError('SYNTHETIC_PRIVATE_DETAIL')
 monkeypatch.setattr(f,'request',offline);f.run_once()
 request={'operation':'job','grant_id':gid,'job_id':job['job_id']};result=f.call(request,os.getuid())
 assert result['state']=='RETRY_WAIT' and result['attempts']==1 and result['error']=='SERVICE_UNAVAILABLE'
 assert partial.read_bytes()==original[:10]
 with f.db() as db:
  result.update(state='RUNNING');f.save(db,result)
 f=durable_factory(f.config)
 assert f.call(request,os.getuid())['state']=='QUEUED'
 f.run_once();result=f.call(request,os.getuid());assert result['state']=='SUCCEEDED' and 'error' not in result


def test_queue_recovers_batch_crash_before_expansion(durable_factory,gateway):
 gid=gateway[2];f=durable_factory();job=f.call({'operation':'batch','grant_id':gid},os.getuid())
 with f.db() as db:job['state']='RUNNING';f.save(db,job)
 f=durable_factory(f.config);assert f.run_once()
 with f.db() as db:assert db.execute("SELECT COUNT(*) FROM jobs WHERE kind='artifact'").fetchone()[0]==4


def test_queue_revoke_blocks_and_does_not_renew_around_denial(durable_factory,gateway):
 _,m,gid,policy,g,_=gateway;enable_long_grant(gateway);f=durable_factory()
 job=f.call({'operation':'batch','grant_id':gid},os.getuid());policy.write_text(json.dumps({'grants':{gid:dict(g,status='REVOKED')}}))
 f.run_once();request={'operation':'job','grant_id':gid,'job_id':job['job_id']}
 assert f.call(request,os.getuid())['state']=='BLOCKED'
 assert not f.run_once()
 # Resume is only a local queue action; it cannot override upstream denial.
 f.call(dict(request,operation='resume'),os.getuid());f.run_once()
 assert f.call(request,os.getuid())['state']=='BLOCKED'


def test_queue_disk_pressure_waits_without_deleting_partial(durable_factory,gateway,monkeypatch):
 import errno
 _,m,gid,*_=gateway;f=durable_factory();a=m['artifacts'][0]
 job=f.call({'operation':'download','grant_id':gid,'artifact_id':a['artifact_id']},os.getuid())
 def full(*args):raise OSError(errno.ENOSPC,'synthetic')
 monkeypatch.setattr(f,'check_space',full);f.run_once()
 result=f.call({'operation':'job','grant_id':gid,'job_id':job['job_id']},os.getuid())
 assert result['state']=='RETRY_WAIT' and result['error']=='DISK_PRESSURE' and result['next_at']>time.time()+290


def test_queue_rejects_extra_parameters_and_other_uid(durable_factory,gateway):
 f=durable_factory();request={'operation':'batch','grant_id':gateway[2]}
 for extra in ({'url':'http://127.0.0.1:5005'},{'days':99},{'path':'/tmp/output'}):
  with pytest.raises(ValueError):f.call(dict(request,**extra),os.getuid())
 with pytest.raises(ValueError):f.call(request,os.getuid()+10000)


def test_offline_queue_has_a_hard_local_lifetime(durable_factory,gateway):
 f=durable_factory();gid=gateway[2];job=f.call({'operation':'batch','grant_id':gid},os.getuid())
 with f.db() as db:job['created_at']=time.time()-15*86400;f.save(db,job)
 f.run_once();result=f.call({'operation':'job','grant_id':gid,'job_id':job['job_id']},os.getuid())
 assert result['state']=='BLOCKED' and result['error']=='APPROVAL_EXPIRED'


def test_server_restart_retains_leases_and_total_usage(gateway):
 from src.research_export import leases
 _,m,gid,_,_,store=gateway;g=enable_long_grant(gateway);peer=g['certificate_sha256']
 token=leases.renew(store,gid,peer)['token'];store.charge(gid,peer,123,lease=token)
 restarted=Store(store.policy,store.published,store.state)
 assert leases.renew(restarted,gid,peer)['token']==token
 restarted.charge(gid,peer,1,lease=token)
 with restarted.db() as db:assert db.execute('SELECT bytes FROM usage WHERE grant_id=?',(gid,)).fetchone()[0]==124


def test_expired_lease_refresh_during_stream_preserves_identity(gateway,durable_factory,monkeypatch):
 get,m,gid,policy,g,store=gateway
 a=m['artifacts'][0];path=store.published/m['export_id']/'artifacts'/a['file'];os.chmod(path,0o600);payload=os.urandom(1024*1024);path.write_bytes(payload);os.chmod(path,0o440);a['bytes']=len(payload);a['sha256']=hashlib.sha256(payload).hexdigest()
 manifest=store.published/m['export_id']/'manifest.json';os.chmod(manifest,0o600);manifest.write_text(json.dumps(m));os.chmod(manifest,0o440)
 g=dict(g,not_before=int(time.time()),expires_at=int(time.time())+14*86400,lease_seconds=600,manifest_sha256=digest(manifest),max_bytes=3*1024**2,rate_bytes_s=1024**2);policy.write_text(json.dumps({'grants':{gid:g}}))
 f=durable_factory();real=f.renew_lease
 def frequent(gid):
  if gid in f.leases:f.leases[gid]['refresh_at']=0
  return real(gid)
 monkeypatch.setattr(f,'renew_lease',frequent)
 job=f.call({'operation':'download','grant_id':gid,'artifact_id':a['artifact_id']},os.getuid());f.run_once()
 result=f.call({'operation':'job','grant_id':gid,'job_id':job['job_id']},os.getuid())
 assert result['state']=='SUCCEEDED' and Path(result['path']).read_bytes()==payload
 with store.db() as db:assert db.execute("SELECT COUNT(*) FROM audit WHERE action='lease'").fetchone()[0]>=4


def test_stale_lease_recovers_once_without_grant_replacement(durable_factory,gateway):
 _,m,gid,_,_,store=gateway;enable_long_grant(gateway);f=durable_factory();f.manifest(gid);token=f.leases[gid]['token']
 with store.db() as db:db.execute('UPDATE leases SET until=0 WHERE grant_id=?',(gid,))
 f.manifest(gid);assert f.leases[gid]['token']==token


def test_quota_exhaustion_cannot_be_renewed(gateway):
 from src.research_export import leases
 from src.research_export.gateway import Denied
 _,_,gid,_,_,store=gateway;g=enable_long_grant(gateway);peer=g['certificate_sha256'];token=leases.renew(store,gid,peer)['token']
 store.charge(gid,peer,g['max_bytes'],lease=token)
 with pytest.raises(Denied):leases.renew(store,gid,peer)


def test_lingering_stream_slot_is_retryable_not_authorization_denial(gateway):
 get,m,gid,_,g,store=gateway;peer=g['certificate_sha256'];store.enter(peer)
 try:
  with pytest.raises(urllib.error.HTTPError) as e:get('/v1/artifacts/'+m['artifacts'][0]['artifact_id']+'/content')
  assert e.value.code==429 and e.value.headers['Retry-After']=='5'
 finally:store.leave(peer)
 with get('/v1/artifacts/'+m['artifacts'][0]['artifact_id']+'/content') as r:assert r.status==200


def test_gateway_rebinds_after_transfer_but_rejects_duplicate_listener(gateway,certificates):
 _,m,gid,_,_,store=gateway;p=certificates
 ctx=ssl.SSLContext(ssl.PROTOCOL_TLS_SERVER);ctx.verify_mode=ssl.CERT_REQUIRED;ctx.load_cert_chain(p/'server.crt',p/'server.key');ctx.load_verify_locations(p/'ca.crt')
 s=Server(('127.0.0.1',0),store,ctx);address=s.server_address;t=threading.Thread(target=s.serve_forever,daemon=True);t.start()
 client=ssl.create_default_context(cafile=p/'ca.crt');client.load_cert_chain(p/'client.crt',p/'client.key')
 opener=urllib.request.build_opener(urllib.request.ProxyHandler({}),urllib.request.HTTPSHandler(context=client))
 try:
  with pytest.raises(OSError):Server(address,store,ctx)
  req=urllib.request.Request('https://127.0.0.1:'+str(address[1])+'/v1/artifacts/'+m['artifacts'][0]['artifact_id']+'/content',headers={'X-Export-Grant':gid})
  with opener.open(req,timeout=3) as r:assert r.read()
 finally:s.shutdown();s.server_close();t.join(timeout=3)
 replacement=Server(address,store,ctx);replacement.server_close()
