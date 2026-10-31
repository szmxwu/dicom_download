"""Trusted-side publisher: explicit review, type allowlist, rebuilt metadata."""
import gzip,hashlib,hmac,json,os,re,secrets,shutil,stat,struct,time,zipfile
from pathlib import Path
import numpy as np
import nibabel as nib
from PIL import Image
from openpyxl import load_workbook,Workbook

POLICY='research-v1'
EXTENSIONS=('.nii.gz','.nii','.npz','.png','.xlsx')
NUMERIC_FIELDS={'Rows','Columns','SliceThickness','PixelSpacing','WindowCenter','WindowWidth','KVP','XRayTubeCurrent','ExposureTime','Exposure','InstanceNumber','SeriesNumber','Low_quality','FileIndex','TotalFilesInSeries','FilesReadForMetadata','SliceLocation','ImagePositionPatient','ImageOrientationPatient'}
CATEGORICAL={'Modality':{'CT','MR','DX','CR','DR','MG','OT'},'BodyPartExamined':{'HEAD','CHEST','ABDOMEN','PELVIS','SPINE','KNEE','ANKLE','WRIST','ELBOW','HIP','SHOULDER','FOOT','HAND','SKULL','NECK','LIMB'}}
class Rejected(ValueError):pass

def digest(path):
 h=hashlib.sha256()
 with open(path,'rb') as f:
  for chunk in iter(lambda:f.read(1024*1024),b''):h.update(chunk)
 return h.hexdigest()

def regular(path,root):
 p=Path(path)
 if p.is_absolute() or '..' in p.parts:raise Rejected('unsafe_path')
 current=root
 for part in p.parts:
  current=current/part
  if current.is_symlink():raise Rejected('symlink')
 s=current.stat()
 if not stat.S_ISREG(s.st_mode) or s.st_nlink!=1:raise Rejected('not_single_regular_file')
 if not current.resolve().is_relative_to(root.resolve()):raise Rejected('outside_source')
 return current

def copy_bounded(src,dst,limit=8*1024**3):
 total=0
 for b in iter(lambda:src.read(1024*1024),b''):
  total+=len(b)
  if total>limit:raise Rejected('decoded_size_limit')
  dst.write(b)

def clean_nifti(src,dst):
 opener=gzip.open if src.name.endswith('.gz') else open
 with opener(src,'rb') as f:
  block=f.read(4)
  size=next((struct.unpack(e+'i',block)[0] for e in ('<','>') if struct.unpack(e+'i',block)[0] in (348,540)),None)
  if size is None:raise Rejected('invalid_nifti')
  raw=block+f.read(size-4)
  header=(nib.Nifti1Header if size==348 else nib.Nifti2Header)(binaryblock=raw,check=True)
  if not bytes(header['magic']).startswith(b'n+'):raise Rejected('paired_nifti_unsupported')
  offset=int(header['vox_offset'])
  if not size+4<=offset<=1024*1024:raise Rejected('invalid_nifti_offset')
  for field in header.structarr.dtype.names:
   if header.structarr.dtype[field].kind in ('S','U') and field!='magic':header[field]=b''
  f.read(offset-size)
  if dst.name.endswith('.gz'):
   stream=gzip.GzipFile(filename='',mode='wb',fileobj=dst.open('wb'),mtime=0)
  else:stream=dst.open('wb')
  with stream as out:
   out.write(header.binaryblock);out.write(bytes(offset-size))
   expected=int(np.prod(header.get_data_shape(),dtype=object))*header.get_data_dtype().itemsize
   if not 0<expected<=8*1024**3:raise Rejected('nifti_payload_limit')
   dtype=header.get_data_dtype()
   if dtype.kind not in 'biufc':raise Rejected('unsupported_nifti_pixel_type')
   remaining=expected;first_value=None;different=False
   while remaining:
    chunk=f.read(min(remaining,1024*1024))
    if not chunk:raise Rejected('truncated_nifti')
    values=np.frombuffer(chunk,dtype=dtype)
    if not np.isfinite(values).all():raise Rejected('nonfinite_nifti')
    if first_value is None:first_value=values[0]
    different=different or bool(np.any(values!=first_value))
    out.write(chunk);remaining-=len(chunk)
   if not different:raise Rejected('constant_nifti')
   if f.read(1):raise Rejected('unexpected_nifti_trailing_bytes')
 # Validation reads image shape/header; raw payload bytes and geometry were preserved.
 image=nib.load(str(dst))
 if not all(x>0 for x in image.shape):raise Rejected('invalid_image_shape')

def clean_npz(src,dst):
 allowed={'data','image','volume','spacing','origin','direction','affine'}
 with zipfile.ZipFile(src) as z:
  if sum(x.file_size for x in z.infolist())>2*1024**3 or len(z.infolist())>16:raise Rejected('npz_limits')
  total_array_bytes=0
  for member in z.infolist():
   if member.filename not in {name+'.npy' for name in allowed}:raise Rejected('npz_keys')
   with z.open(member) as stream:
    version=np.lib.format.read_magic(stream)
    if version==(1,0):shape,order,dtype=np.lib.format.read_array_header_1_0(stream)
    elif version==(2,0):shape,order,dtype=np.lib.format.read_array_header_2_0(stream)
    else:raise Rejected('npz_header_version')
    if dtype.kind not in 'biufc' or dtype.hasobject or not 0<=len(shape)<=5 or any(not isinstance(n,int) or n<1 for n in shape):raise Rejected('npz_array_header')
    size=int(np.prod(shape,dtype=object))*dtype.itemsize
    total_array_bytes+=size
    if total_array_bytes>2*1024**3 or stream.tell()+size!=member.file_size:raise Rejected('npz_shape_size_limit')
 with np.load(src,allow_pickle=False) as data:
  if not set(data.files)<=allowed or not set(data.files)&{'data','image','volume'}:raise Rejected('npz_keys')
  arrays={}
  for name in data.files:
   value=data[name]
   if value.dtype.kind not in 'biufc' or not np.isfinite(value).all():raise Rejected('npz_non_numeric_or_nonfinite')
   if name in {'data','image','volume'} and (value.size<2 or not np.any(value!=value.flat[0])):raise Rejected('constant_npz_image')
   arrays[name]=value
  np.savez_compressed(dst,**arrays)

def clean_excel(src,dst,subject):
 with zipfile.ZipFile(src) as z:
  if sum(x.file_size for x in z.infolist())>64*1024**2:raise Rejected('xlsx_size_limit')
 book=load_workbook(src,read_only=True,data_only=False)
 out=Workbook();out.remove(out.active)
 identities={'PatientID':set(),'AccessionNumber':set()}
 try:
  for index,sheet in enumerate(book.worksheets):
   if sheet.sheet_state!='visible':continue
   rows=sheet.iter_rows(values_only=True);head=next(rows,())
   identity_indices=[(i,k) for i,k in enumerate(head) if k in identities]
   keep=[(i,str(k)) for i,k in enumerate(head) if k in NUMERIC_FIELDS or k in CATEGORICAL]
   target=out.create_sheet('Metadata_'+str(index+1));target.append(['ResearchID']+[k for _,k in keep])
   for number,row in enumerate(rows):
    if number>=100000:raise Rejected('xlsx_row_limit')
    for i,key in identity_indices:
     if i<len(row) and row[i] is not None:identities[key].add(str(row[i]))
     if len(identities[key])>1:raise Rejected('one_case_per_excel_required')
    values=[subject]
    for i,key in keep:
     v=row[i] if i<len(row) else None
     if key in CATEGORICAL:v=v if isinstance(v,str) and v in CATEGORICAL[key] else None
     elif v is not None:
      if isinstance(v,(int,float)) and not isinstance(v,bool):pass
      elif isinstance(v,str) and re.fullmatch(r'[\d+.,eE\-\\\[\] ]{1,160}',v):pass
      else:v=None
     values.append(v)
    target.append(values)
  if not out.worksheets:raise Rejected('xlsx_no_visible_sheet')
  out.save(dst)
 finally:book.close();out.close()

def clean_png(src,dst):
 with Image.open(src) as image:
  if image.width*image.height>64_000_000 or getattr(image,'n_frames',1)!=1:raise Rejected('png_dimensions')
  if image.format!='PNG' or image.mode not in ('L','RGB','RGBA','I;16'):raise Rejected('png_mode')
  # Preserve numerical pixels/mode; a new image has no inherited text/EXIF/profile.
  pixels=np.asarray(image)
  if pixels.size<2 or not np.any(pixels!=pixels[0,0]):raise Rejected('constant_preview_requires_review')
  Image.fromarray(pixels).save(dst,format='PNG')

def publish(source,review,root,project,key):
 source=Path(source).resolve();root=Path(root);files=review.get('files',[])
 if review.get('policy_version')!=POLICY or not review.get('reviewer') or not 1<=len(files)<=1000:raise Rejected('review_required')
 if not re.fullmatch(r'[a-zA-Z0-9_-]{1,64}',project):raise Rejected('invalid_project')
 export=secrets.token_hex(16);stage=root/'staging'/export;stage.mkdir(parents=True,mode=0o700);artifacts=[];subjects=set()
 try:
  (stage/'artifacts').mkdir(mode=0o750)
  output_bytes=0
  for entry in files:
   if shutil.disk_usage(root).free<24*1024**3:raise Rejected('publication_capacity_reserve')
   src=regular(entry['path'],source);before=src.stat();suffix=next((s for s in EXTENSIONS if src.name.endswith(s)),None)
   if not suffix:raise Rejected('unsupported_file_type')
   sha=digest(src)
   if sha!=entry.get('sha256'):raise Rejected('source_digest_mismatch')
   if suffix!='.xlsx' and (entry.get('pixel_review',{}).get('status')!='approved' or not entry['pixel_review'].get('evidence')):raise Rejected('pixel_review_required')
   if suffix!='.xlsx' and (entry.get('quality_review',{}).get('status')!='approved' or not entry['quality_review'].get('evidence')):raise Rejected('quality_review_required')
   case=entry.get('case_ref')
   if not isinstance(case,str) or not case or len(case)>256:raise Rejected('case_reference_required')
   subject=hmac.new(key,(project+'\0'+case).encode(),hashlib.sha256).hexdigest()[:32];subjects.add(subject)
   series=entry.get('series_number')
   if series is not None and (not isinstance(series,int) or isinstance(series,bool) or not 0<=series<=1000000):raise Rejected('invalid_series_number')
   artifact=secrets.token_hex(16);path=stage/'artifacts'/(artifact+suffix)
   if suffix.startswith('.nii'):clean_nifti(src,path)
   elif suffix=='.npz':clean_npz(src,path)
   elif suffix=='.png':clean_png(src,path)
   else:clean_excel(src,path,subject)
   after=src.stat()
   if (before.st_ino,before.st_size,before.st_mtime_ns)!=(after.st_ino,after.st_size,after.st_mtime_ns) or digest(src)!=sha:raise Rejected('source_changed')
   output_bytes+=path.stat().st_size
   if output_bytes>20*1024**3:raise Rejected('initial_publication_size_limit')
   os.chmod(path,0o440)
   artifacts.append({'artifact_id':artifact,'name':subject+'_'+artifact[:8]+suffix,'file':artifact+suffix,'format':suffix,'research_id':subject,'series_number':entry.get('series_number'),'bytes':path.stat().st_size,'sha256':digest(path)})
  manifest={'export_id':export,'project_id':project,'version':1,'policy_version':POLICY,'publisher_sha256':digest(Path(__file__)),'case_count':len(subjects),'created_at':int(time.time()),'artifacts':artifacts}
  (stage/'manifest.json').write_text(json.dumps(manifest,sort_keys=True)+'\n');os.chmod(stage/'manifest.json',0o440)
  target=root/'published'/export;target.parent.mkdir(exist_ok=True,mode=0o750)
  os.chmod(stage,0o750);os.rename(stage,target)
  return manifest
 except BaseException:
  quarantine=root/'quarantine';quarantine.mkdir(exist_ok=True,mode=0o700);os.rename(stage,quarantine/export);raise
